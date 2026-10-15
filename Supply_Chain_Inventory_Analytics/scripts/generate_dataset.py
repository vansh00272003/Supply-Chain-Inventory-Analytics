import os
import random
import datetime
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Set seeds for reproducibility
random.seed(42)
np.random.seed(42)

print("Starting synthetic dataset generation for Supply Chain & Inventory Analytics...")

# ---------------------------------------------------------
# 1. PRODUCTS TABLE (240 Products)
# ---------------------------------------------------------
departments_subcategories = {
    "Grocery": [
        ("Rice & Grains", ["Tata Daily Rice 5kg", "India Gate Basmati Rice 5kg", "Fortune Biryani Rice 1kg", "Daawat Rozana Rice 5kg", "Organic Tattva Sona Masoori 5kg"]),
        ("Flour & Staples", ["Aashirvaad Select Wheat Flour 5kg", "Pillsbury Chakki Atta 5kg", "Fortune Besan 1kg", "Fortune Maida 1kg", "Tata Sampann Suji 1kg"]),
        ("Pulses", ["Tata Sampann Toor Dal 1kg", "Tata Sampann Moong Dal 1kg", "Catch Chana Dal 1kg", "Organic Tattva Urad Dal 1kg", "Fortune Rajma 1kg"]),
        ("Cooking Oils", ["Fortune Cooking Oil 1L", "Saffola Gold Oil 1L", "Dhara Mustard Oil 1L", "Gemini Sunflower Oil 1L", "Oleev Olive Oil 1L"]),
        ("Spices", ["Catch Red Chilli Powder 200g", "MDH Garam Masala 100g", "Everest Turmeric Powder 200g", "Tata Sampann Coriander Powder 200g", "Catch Cumin Seeds 100g"]),
        ("Instant Food", ["Nestle Instant Noodles 70g", "Knorr Soups Classic 45g", "MTR Ready to Eat Curry 300g", "Yippee Noodles Pack 280g", "Bambino Roasted Vermicelli 400g"]),
        ("Sauces & Spreads", ["Kissan Tomato Ketchup 1kg", "Maggi Hot & Sweet Sauce 500g", "FunFoods Mayonnaise 250g", "Nutella Hazelnut Spread 350g", "Dr. Oetker Peanut Butter 340g"]),
        ("Canned Food", ["Del Monte Sweet Corn 400g", "Heinz Baked Beans 415g", "Pineapple Slices Canned 450g", "Mushroom Buttons Canned 400g"])
    ],
    "Beverages": [
        ("Soft Drinks", ["Coca Cola 1.5L", "Pepsi 1.5L", "Thums Up 1.5L", "Sprite 1.5L", "Seven Up 1.5L"]),
        ("Juices", ["Tropicana Orange Juice 1L", "Real Mixed Fruit Juice 1L", "B Natural Mango Juice 1L", "Paper Boat Anardana 250ml", "Raw Pressery Coconut Water 200ml"]),
        ("Tea & Coffee", ["Tata Tea Gold 500g", "Brooke Bond Red Label 500g", "Nescafe Classic Coffee 100g", "Bru Instant Coffee 100g", "Organic India Green Tea 25s"]),
        ("Energy Drinks", ["Red Bull Energy Drink 250ml", "Monster Energy 350ml", "Sting Energy Drink 250ml"]),
        ("Water", ["Bisleri Mineral Water 1L", "Kinley Mineral Water 1L", "Himalayan Mineral Water 1L"])
    ],
    "Snacks": [
        ("Biscuits", ["Britannia Family Pack Biscuits 500g", "Parle-G Gold 1kg", "Sunfeast Dark Fantasy 300g", "Monaco Salted Biscuits 200g", "Oreo Chocolate Biscuits 120g"]),
        ("Chips", ["Lays Classic Chips 100g", "Lays Magic Masala 100g", "Kurkure Masala Munch 90g", "Pringles Original 107g", "Bingo Mad Angles 80g"]),
        ("Namkeen", ["Haldirams Bhujia Sev 400g", "Haldirams Khatta Meetha 350g", "Bikanervala Navratan Mix 400g", "Haldirams Aloo Bhujia 400g"]),
        ("Chocolates", ["Cadbury Dairy Milk Silk 150g", "Nestle KitKat 4-Finger 38g", "Ferrero Rocher 16s", "Snickers Peanut Bar 50g", "Amul Dark Chocolate 150g"]),
        ("Nuts", ["Happilo Premium Almonds 250g", "Happilo Cashews 250g", "Nutraj Raisins 250g", "Solimo Walnuts 200g"])
    ],
    "Dairy": [
        ("Milk", ["Amul Premium Milk 1L", "Mother Dairy Toned Milk 1L", "Nandini Milk 1L", "Country Delight Pure Milk 1L"]),
        ("Butter & Spreads", ["Amul Pasteurised Butter 500g", "Mother Dairy Butter 500g", "Nutralite Garlic Spread 200g", "Amul Delicious Spread 200g"]),
        ("Cheese", ["Amul Cheese Blocks 200g", "Britannia Cheese Slices 200g", "Mother Dairy Processed Cheese 200g", "Go Cheese Spread 200g"]),
        ("Yogurt", ["Amul Masti Dahi 400g", "Mother Dairy Ultimate Dahi 400g", "Epigamia Greek Yogurt 90g", "Milky Mist Curd 500g"]),
        ("Eggs", ["Farm Fresh Eggs 6s", "Farm Fresh Eggs 12s", "Eggoz Organic Eggs 6s"])
    ],
    "Bakery": [
        ("Bread", ["Britannia Whole Wheat Bread 400g", "Modern Brown Bread 400g", "English Oven Garlic Bread 250g", "Harvest Gold Sandwich Bread 400g"]),
        ("Cakes & Pastries", ["Britannia Chocolate Cake 250g", "Bonn Fruit Cake 300g", "Winkies Slice Cake 150g", "Elite Plum Cake 400g"]),
        ("Buns & Pav", ["English Oven Burger Buns 4s", "Modern Pav Pack 6s", "Harvest Gold Pizza Base 2s"])
    ],
    "Fresh Produce": [
        ("Fruits", ["Fresh Shimla Apples 1kg", "Fresh Robusta Bananas 1kg", "Nagpur Oranges 1kg", "Thompson Seedless Grapes 500g", "Fresh Pomegranate 1kg"]),
        ("Vegetables", ["Fresh Onions 5kg", "Fresh Potatoes 5kg", "Fresh Tomatoes 1kg", "Fresh Hybrid Carrots 500g", "Fresh Green Capsicum 500g"]),
        ("Herbs", ["Fresh Coriander Leaves 100g", "Fresh Mint Leaves 100g", "Fresh Ginger 250g", "Fresh Garlic 250g"])
    ],
    "Meat & Seafood": [
        ("Chicken", ["Fresh Chicken Curry Cut 500g", "Fresh Chicken Breast Boneless 500g", "Fresh Chicken Drumsticks 500g"]),
        ("Mutton", ["Fresh Mutton Curry Cut 500g", "Fresh Mutton Boneless 500g"]),
        ("Fish & Seafood", ["Fresh Rohu Fish Cut 500g", "Fresh Seer Fish Steaks 500g", "Fresh White Prawns 500g"])
    ],
    "Frozen Foods": [
        ("Frozen Vegetables", ["Safal Green Peas 1kg", "Safal Sweet Corn 500g", "Safal Mixed Veg 500g"]),
        ("Frozen Snacks", ["McCain French Fries 750g", "McCain Aloo Tikki 400g", "Godrej Yummiez Chicken Nuggets 500g", "Sumeru Veg Momos 300g"]),
        ("Ice Cream", ["Amul Vanilla Gold 1L", "Kwality Wall's Chocolate Fudge 700ml", "Baskin Robbins Mississippi Mud 500ml", "Havmor Butterscotch 1L"])
    ],
    "Household": [
        ("Laundry", ["Surf Excel Laundry Powder 2kg", "Ariel Matic Liquid Detergent 1L", "Tide Plus Detergent 2kg", "Comfort Fabric Conditioner 860ml"]),
        ("Cleaning", ["Harpic Disinfectant Toilet Cleaner 1L", "Colin Glass Cleaner 500ml", "Lizol Surface Cleaner 1L", "Vim Dishwash Gel 750ml"]),
        ("Paper Products", ["Origami Tissue Napkins 100s", "Selpak Toilet Paper 4 Rolls", "Kitchen Towel Roll 2s"])
    ],
    "Personal Care": [
        ("Bath & Body", ["Dove Cream Beauty Bar 125g x 3", "Dettol Original Soap 125g x 4", "Pears Soft Soap 125g x 3", "Fiama Gel Bar 125g x 3"]),
        ("Hair Care", ["Head & Shoulders Shampoo 650ml", "Clinic Plus Strong Shampoo 650ml", "Pantene Hairfall Conditioner 200ml", "Parachute Coconut Hair Oil 300ml"]),
        ("Oral Care", ["Colgate Fresh Toothpaste 150g", "Sensodyne Repair Toothpaste 100g", "Oral-B Soft Toothbrush 4s", "Listerine Mouthwash 500ml"]),
        ("Skin Care", ["Nivea Soft Cream 200ml", "Ponds Light Moisturizer 200ml", "Vaseline Intensive Care Lotion 400ml", "Garnier Face Wash 100g"])
    ],
    "Baby Care": [
        ("Diapers", ["Pampers Baby Diapers 24s", "Huggies Wonder Pants 30s", "MamyPoko Pants 32s"]),
        ("Baby Food", ["Nestle Cerelac Wheat Apple 300g", "Nestle Lactogen 1 400g", "Slurrp Farm Organic Porridge 200g"]),
        ("Baby Bath & Wipes", ["Johnson's Baby Shampoo 500ml", "Himalaya Baby Wipes 72s", "Mother Sparsh Plant Wipes 72s"])
    ],
    "Pet Supplies": [
        ("Dog Food", ["Pedigree Dog Food 1kg", "Pedigree Adult Chicken 3kg", "Drools Chicken & Egg 3kg", "Purepet Biscuit Dog Treats 500g"]),
        ("Cat Food", ["Whiskas Cat Food Tuna 1.2kg", "Me-O Cat Food Salmon 1.1kg", "Sheba Premium Cat Food 70g"]),
        ("Pet Care", ["Choostix Dog Chew Bones 250g", "Pet Hygiene Shampoo 200ml"])
    ],
    "Health & Wellness": [
        ("Vitamins", ["Revital H Daily Health Capsules 30s", "Limcee Vitamin C 15s", "Becosules Z Capsules 20s"]),
        ("First Aid", ["Dettol Antiseptic Liquid 550ml", "Hansaplast Bandage 100s", "Iodex Fast Relief Balm 40g", "Volini Pain Relief Spray 100g"]),
        ("Personal Health", ["Chyawanprash Dabur 1kg", "Horlicks Health Drink 500g", "Bournvita Health Drink 500g"])
    ],
    "Stationery": [
        ("School Supplies", ["Classmate Notebook A4 300 Pages", "Classmate Drawing Book", "Faber Castell Colored Pencils 24s"]),
        ("Office Supplies", ["JK Copier A4 Paper Rim 500 Sheets", "Cello Ball Pens Pack of 10", "Camel Glue Stick 15g"]),
        ("Writing", ["Parker Vector Stainless Pen", "Uniball Eye Rollerball Pen 3s"])
    ],
    "Electronics": [
        ("Batteries", ["Eveready AA Batteries 4s", "Duracell Ultra AAA Batteries 4s", "Panasonic Rechargeable Batteries 2s"]),
        ("Mobile Accessories", ["Boat Micro USB Cable 1m", "Mi 20W Fast Charger Adapter", "SanDisk 64GB MicroSD Card"]),
        ("Small Electronics", ["Philips LED Bulb 9W 2s", "Havells Extension Board 4 Socket", "Syska Emergency Lantern"])
    ],
    "Home & Kitchen": [
        ("Cookware", ["Prestige Cookware 2L", "Hawkins Pressure Cooker 3L", "Wonderchef Non-Stick Fry Pan 24cm"]),
        ("Storage", ["Milton Thermosteel Water Bottle 750ml", "Cello Glass Storage Containers 3s", "Tupperware Executive Lunch Set"]),
        ("Home Accessories", ["Bombay Dyeing Cotton Towel", "Spaces Microfiber Bed Sheet Double"])
    ],
    "Apparel": [
        ("T-Shirts", ["Jockey Cotton Crew T-Shirt M", "Jockey Cotton Crew T-Shirt L", "Monte Carlo Polo T-Shirt L"]),
        ("Socks", ["Jockey Ankle Socks Pack of 3", "Nike Cotton Sports Socks Pack of 3"]),
        ("Innerwear", ["Jockey Mens Briefs Pack of 2", "Lux Cozi Vests Pack of 3"])
    ],
    "Automotive": [
        ("Car Care", ["3M Car Wash Shampoo 500ml", "Turtle Wax Car Polish 250ml", "Areon Car Air Freshener Gel"]),
        ("Bike Care", ["Motul Chain Lube 400ml", "Formula 1 Bike Clean 400ml"]),
        ("Accessories", ["Microfiber Cloth for Car Pack of 4", "Digital Tyre Pressure Gauge"])
    ]
}

brands_list = [
    "Amul", "Tata", "Nestle", "Britannia", "Haldirams", "Aashirvaad", "Fortune", "Surf Excel", 
    "Harpic", "Dettol", "Dove", "Colgate", "Pepsi", "Coca Cola", "Thums Up", "Parle", 
    "Lays", "Kurkure", "Kelloggs", "Pampers", "Whiskas", "Pedigree", "Prestige", "Eveready", 
    "Samsung", "Philips", "Classmate", "Milton", "Nivea", "Dabur", "McCain", "Jockey", 
    "Saffola", "Horlicks", "Happilo", "Safal", "Hawkins", "Huggies", "Ariel", "Lizol"
]

products_data = []
p_counter = 1

# Base prices by department
dept_price_ranges = {
    "Grocery": (30, 450),
    "Beverages": (20, 350),
    "Snacks": (15, 400),
    "Dairy": (25, 250),
    "Bakery": (25, 180),
    "Fresh Produce": (15, 180),
    "Meat & Seafood": (150, 650),
    "Frozen Foods": (80, 380),
    "Household": (40, 550),
    "Personal Care": (35, 480),
    "Baby Care": (120, 850),
    "Pet Supplies": (90, 950),
    "Health & Wellness": (45, 600),
    "Stationery": (25, 450),
    "Electronics": (80, 1500),
    "Home & Kitchen": (150, 2200),
    "Apparel": (120, 850),
    "Automotive": (90, 750)
}

dept_shelf_life = {
    "Grocery": (180, 365),
    "Beverages": (90, 270),
    "Snacks": (90, 180),
    "Dairy": (5, 30),
    "Bakery": (3, 10),
    "Fresh Produce": (3, 14),
    "Meat & Seafood": (2, 7),
    "Frozen Foods": (90, 270),
    "Household": (365, 730),
    "Personal Care": (365, 730),
    "Baby Care": (180, 540),
    "Pet Supplies": (180, 365),
    "Health & Wellness": (180, 730),
    "Stationery": (365, 1095),
    "Electronics": (365, 1095),
    "Home & Kitchen": (365, 1460),
    "Apparel": (365, 1095),
    "Automotive": (365, 730)
}

# Collect all item templates
all_product_templates = []
for dept, subcat_list in departments_subcategories.items():
    for subcat, item_names in subcat_list:
        for name in item_names:
            all_product_templates.append((dept, subcat, name))

# Ensure exactly 240 products by padding or sampling
while len(all_product_templates) < 240:
    # Pick a random template and create a variant
    dept, subcat, name = random.choice(all_product_templates)
    variant_name = f"{name} (Value Pack)"
    all_product_templates.append((dept, subcat, variant_name))

all_product_templates = all_product_templates[:240]

for idx, (dept, subcat, p_name) in enumerate(all_product_templates, start=1):
    pid = f"P{idx:04d}"
    
    # Extract or infer brand
    matched_brand = "Generic"
    for b in brands_list:
        if b.lower() in p_name.lower():
            matched_brand = b
            break
    if matched_brand == "Generic":
        matched_brand = random.choice(brands_list)
        
    min_p, max_p = dept_price_ranges.get(dept, (50, 500))
    unit_cost = round(random.uniform(min_p, max_p), 2)
    # 20% to 45% margin markup
    markup = random.uniform(1.20, 1.45)
    selling_price = round(unit_cost * markup, 2)
    
    reorder_level = random.randint(30, 250)
    safety_stock = int(reorder_level * random.uniform(0.3, 0.5))
    
    min_life, max_life = dept_shelf_life.get(dept, (30, 365))
    shelf_life_days = random.randint(min_life, max_life)
    
    products_data.append({
        "Product_ID": pid,
        "Product_Name": p_name,
        "Department": dept,
        "Subcategory": subcat,
        "Brand": matched_brand,
        "Unit_Cost": unit_cost,
        "Selling_Price": selling_price,
        "Reorder_Level": reorder_level,
        "Safety_Stock": safety_stock,
        "Shelf_Life_Days": shelf_life_days
    })

df_products = pd.DataFrame(products_data)
print(f"Generated {len(df_products)} products.")

# ---------------------------------------------------------
# 2. SUPPLIERS TABLE (20 Suppliers)
# ---------------------------------------------------------
suppliers_info = [
    ("SUP01", "Metro Wholesale", "Chennai", "Tamil Nadu", 3, 4.8),
    ("SUP02", "Reliance Distribution", "Mumbai", "Maharashtra", 4, 4.6),
    ("SUP03", "Shree FMCG Distributors", "Bengaluru", "Karnataka", 5, 4.2),
    ("SUP04", "FreshMart Wholesale", "Hyderabad", "Telangana", 2, 4.9),
    ("SUP05", "South India FMCG", "Chennai", "Tamil Nadu", 4, 3.9),
    ("SUP06", "National Consumer Products", "Pune", "Maharashtra", 6, 4.1),
    ("SUP07", "GreenBasket Foods", "Coimbatore", "Tamil Nadu", 3, 4.7),
    ("SUP08", "Prime Distributors", "Bengaluru", "Karnataka", 5, 3.8),
    ("SUP09", "DailyNeeds Wholesale", "Hyderabad", "Telangana", 4, 4.4),
    ("SUP10", "Global FMCG Traders", "Mumbai", "Maharashtra", 7, 3.5),
    ("SUP11", "Chennai Cash & Carry", "Chennai", "Tamil Nadu", 2, 4.8),
    ("SUP12", "Kaveri Suppliers", "Bengaluru", "Karnataka", 5, 3.7),
    ("SUP13", "Sunrise Distributors", "Pune", "Maharashtra", 4, 4.3),
    ("SUP14", "BigBasket Wholesale", "Bengaluru", "Karnataka", 3, 4.7),
    ("SUP15", "RetailLink Supply", "Hyderabad", "Telangana", 6, 3.4),
    ("SUP16", "Evergreen Foods", "Coimbatore", "Tamil Nadu", 3, 4.5),
    ("SUP17", "Southern Fresh Supply", "Chennai", "Tamil Nadu", 2, 4.6),
    ("SUP18", "ValueMart Distributors", "Mumbai", "Maharashtra", 5, 3.9),
    ("SUP19", "City Wholesale Hub", "Pune", "Maharashtra", 4, 4.2),
    ("SUP20", "Universal Traders", "Hyderabad", "Telangana", 7, 3.2)
]

suppliers_data = []
for sid, sname, city, state, lead, rating in suppliers_info:
    suppliers_data.append({
        "Supplier_ID": sid,
        "Supplier_Name": sname,
        "City": city,
        "State": state,
        "Standard_Lead_Time_Days": lead,
        "Supplier_Rating": rating
    })

df_suppliers = pd.DataFrame(suppliers_data)
print(f"Generated {len(df_suppliers)} suppliers.")

# Map departments to primary supplier candidates
dept_supplier_map = {}
supplier_ids = df_suppliers["Supplier_ID"].tolist()
for dept in departments_subcategories.keys():
    dept_supplier_map[dept] = random.sample(supplier_ids, 5)

# ---------------------------------------------------------
# 3. STORES TABLE (8 Stores)
# ---------------------------------------------------------
stores_info = [
    ("ST01", "Walmart Supercenter Chennai", "Chennai", "Tamil Nadu", "South", 450000),
    ("ST02", "Walmart Supercenter Tambaram", "Tambaram", "Tamil Nadu", "South", 300000),
    ("ST03", "Walmart Supercenter Coimbatore", "Coimbatore", "Tamil Nadu", "South", 250000),
    ("ST04", "Walmart Supercenter Bengaluru", "Bengaluru", "Karnataka", "South", 500000),
    ("ST05", "Walmart Supercenter Hyderabad", "Hyderabad", "Telangana", "South", 400000),
    ("ST06", "Walmart Supercenter Pune", "Pune", "Maharashtra", "West", 350000),
    ("ST07", "Walmart Supercenter Mumbai", "Mumbai", "Maharashtra", "West", 480000),
    ("ST08", "Walmart Distribution Center Chennai", "Chennai", "Tamil Nadu", "South", 650000)
]

stores_data = []
for st_id, sname, city, state, region, cap in stores_info:
    stores_data.append({
        "Store_ID": st_id,
        "Store_Name": sname,
        "City": city,
        "State": state,
        "Region": region,
        "Capacity_Units": cap
    })

df_stores = pd.DataFrame(stores_data)
print(f"Generated {len(df_stores)} stores.")

# ---------------------------------------------------------
# 4. CUSTOMERS TABLE (2,000 Customers)
# ---------------------------------------------------------
cities_pool = ["Chennai", "Bengaluru", "Hyderabad", "Mumbai", "Pune", "Coimbatore", "Tambaram"]
cities_weights = [0.25, 0.22, 0.18, 0.15, 0.10, 0.06, 0.04]
age_groups = ["18-25", "26-35", "36-45", "46-55", "56+"]
age_weights = [0.15, 0.40, 0.25, 0.12, 0.08]
genders = ["Male", "Female"]
segments = ["Regular", "Premium", "Occasional"]
segment_weights = [0.60, 0.15, 0.25]

customers_data = []
for idx in range(1, 2001):
    cid = f"CUST{idx:04d}"
    city = np.random.choice(cities_pool, p=cities_weights)
    age = np.random.choice(age_groups, p=age_weights)
    gender = random.choice(genders)
    seg = np.random.choice(segments, p=segment_weights)
    customers_data.append({
        "Customer_ID": cid,
        "City": city,
        "Age_Group": age,
        "Gender": gender,
        "Customer_Segment": seg
    })

df_customers = pd.DataFrame(customers_data)
print(f"Generated {len(df_customers)} customers.")

# ---------------------------------------------------------
# 5. INVENTORY TABLE (~115,200 Records)
# 240 products x 8 stores x 60 days
# ---------------------------------------------------------
start_date = datetime.date(2026, 8, 1)
num_days = 60
dates_list = [start_date + datetime.timedelta(days=i) for i in range(num_days)]

inventory_data = []
inv_counter = 1

# Product-store base stock levels
prod_store_stock = {}
for p in products_data:
    pid = p["Product_ID"]
    reorder = p["Reorder_Level"]
    prod_store_stock[pid] = {}
    for st in df_stores["Store_ID"]:
        # DC ST08 holds higher stock
        mult = 2.5 if st == "ST08" else random.uniform(0.6, 1.8)
        prod_store_stock[pid][st] = int(reorder * mult)

for d in dates_list:
    d_str = d.strftime("%Y-%m-%d")
    is_weekend = d.weekday() in [5, 6]
    
    for p in products_data:
        pid = p["Product_ID"]
        reorder = p["Reorder_Level"]
        safety = p["Safety_Stock"]
        
        for st in df_stores["Store_ID"]:
            inv_id = f"INV{inv_counter:06d}"
            inv_counter += 1
            
            curr_base = prod_store_stock[pid][st]
            
            # Daily inventory fluctuation (sales draw-down, periodic reorders)
            drawdown = random.randint(2, 25) if is_weekend else random.randint(1, 15)
            curr_base -= drawdown
            
            # Stockout condition (3.5% chance or if depleted)
            if curr_base <= 0 or random.random() < 0.035:
                stock_on_hand = 0
                reserved_stock = 0
                curr_base = 0
            else:
                stock_on_hand = max(0, curr_base)
                reserved_stock = random.randint(0, min(15, stock_on_hand))
                
            available_stock = stock_on_hand - reserved_stock
            
            # Replenishment trigger if stock is low
            if stock_on_hand <= safety and random.random() < 0.65:
                # Add replenishment shipment
                curr_base += int(reorder * random.uniform(1.2, 2.0))
                
            prod_store_stock[pid][st] = curr_base
            
            inventory_data.append({
                "Inventory_ID": inv_id,
                "Inventory_Date": d_str,
                "Product_ID": pid,
                "Store_ID": st,
                "Stock_On_Hand": stock_on_hand,
                "Reserved_Stock": reserved_stock,
                "Available_Stock": available_stock
            })

df_inventory = pd.DataFrame(inventory_data)
print(f"Generated {len(df_inventory)} inventory records.")

# ---------------------------------------------------------
# 6. SALES ORDERS TABLE (15,000+ Records)
# ---------------------------------------------------------
sales_data = []
ord_counter = 1

payment_methods = ["Cash", "UPI", "Credit Card", "Debit Card", "Gift Card"]
payment_weights = [0.15, 0.45, 0.25, 0.12, 0.03]
discounts = [0.0, 0.02, 0.05, 0.10, 0.15]
discount_weights = [0.55, 0.15, 0.15, 0.10, 0.05]

# Weight stores for sales
store_sales_weights = [0.18, 0.10, 0.08, 0.22, 0.16, 0.10, 0.14, 0.02] # ST08 is DC, low retail sales
store_ids_list = df_stores["Store_ID"].tolist()

# Map products by price for quantity rules
prod_dict = {p["Product_ID"]: p for p in products_data}
prod_ids = list(prod_dict.keys())

num_sales_orders = 15200
print(f"Generating {num_sales_orders} sales order records...")

for i in range(num_sales_orders):
    oid = f"ORD{ord_counter:05d}"
    ord_counter += 1
    
    # Pick date
    o_date = random.choice(dates_list)
    d_str = o_date.strftime("%Y-%m-%d")
    
    cust_id = random.choice(customers_data)["Customer_ID"]
    pid = random.choice(prod_ids)
    p_info = prod_dict[pid]
    
    st_id = np.random.choice(store_ids_list, p=store_sales_weights)
    
    # Quantity based on unit selling price
    sp = p_info["Selling_Price"]
    uc = p_info["Unit_Cost"]
    
    if sp > 1000:
        qty = random.randint(1, 2)
    elif sp > 300:
        qty = random.randint(1, 4)
    else:
        qty = random.randint(1, 8)
        
    disc = np.random.choice(discounts, p=discount_weights)
    
    sales_amt = round(qty * sp * (1.0 - disc), 2)
    cogs = round(qty * uc, 2)
    gross_profit = round(sales_amt - cogs, 2)
    
    pm = np.random.choice(payment_methods, p=payment_weights)
    
    sales_data.append({
        "Order_ID": oid,
        "Order_Date": d_str,
        "Customer_ID": cust_id,
        "Product_ID": pid,
        "Store_ID": st_id,
        "Quantity": qty,
        "Discount_Pct": disc,
        "Sales_Amount": sales_amt,
        "Payment_Method": pm,
        "Unit_Selling_Price": sp,
        "COGS": cogs,
        "Gross_Profit": gross_profit
    })

df_sales = pd.DataFrame(sales_data)
print(f"Generated {len(df_sales)} sales orders.")

# ---------------------------------------------------------
# 7. PURCHASES TABLE (3,000+ Purchase Orders)
# ---------------------------------------------------------
purchases_data = []
po_counter = 1
num_pos = 3100
print(f"Generating {num_pos} purchase orders...")

# Assign supplier rating lookup for delay probability
sup_rating_dict = {s["Supplier_ID"]: s for s in suppliers_data}

for i in range(num_pos):
    poid = f"PO{po_counter:04d}"
    po_counter += 1
    
    pid = random.choice(prod_ids)
    p_info = prod_dict[pid]
    dept = p_info["Department"]
    
    # Select supplier from department candidates
    cand_sups = dept_supplier_map.get(dept, supplier_ids)
    sup_id = random.choice(cand_sups)
    s_info = sup_rating_dict[sup_id]
    
    st_id = random.choice(store_ids_list)
    
    ord_date = random.choice(dates_list[:50]) # ensure expected date fits in horizon
    std_lead = s_info["Standard_Lead_Time_Days"]
    rating = s_info["Supplier_Rating"]
    
    exp_date = ord_date + datetime.timedelta(days=std_lead)
    
    # Delay probability higher for low rated suppliers
    if rating < 3.8:
        delay_days = random.choices([-1, 0, 1, 2, 3, 5, 7], weights=[0.05, 0.35, 0.20, 0.15, 0.10, 0.10, 0.05])[0]
    elif rating < 4.5:
        delay_days = random.choices([-1, 0, 1, 2, 3], weights=[0.10, 0.65, 0.15, 0.07, 0.03])[0]
    else:
        delay_days = random.choices([-1, 0, 1], weights=[0.15, 0.80, 0.05])[0]
        
    rec_date = exp_date + datetime.timedelta(days=delay_days)
    lead_time_days = (rec_date - ord_date).days
    on_time = "Yes" if delay_days <= 0 else "No"
    
    qty_ordered = random.randint(100, 1200)
    purch_cost = round(qty_ordered * p_info["Unit_Cost"], 2)
    
    purchases_data.append({
        "PO_ID": poid,
        "Order_Date": ord_date.strftime("%Y-%m-%d"),
        "Product_ID": pid,
        "Supplier_ID": sup_id,
        "Store_ID": st_id,
        "Expected_Date": exp_date.strftime("%Y-%m-%d"),
        "Received_Date": rec_date.strftime("%Y-%m-%d"),
        "Quantity_Ordered": qty_ordered,
        "Purchase_Cost": purch_cost,
        "Lead_Time_Days": lead_time_days,
        "Delay_Days": delay_days,
        "On_Time": on_time
    })

df_purchases = pd.DataFrame(purchases_data)
print(f"Generated {len(df_purchases)} purchase orders.")

# ---------------------------------------------------------
# 8. RETURNS TABLE (1,000+ Returns)
# ---------------------------------------------------------
returns_data = []
ret_counter = 1
num_returns = 1050
print(f"Generating {num_returns} return records...")

return_reasons = ["Damaged", "Customer Changed Mind", "Wrong Item", "Quality Issue", "Expired"]
reason_weights = [0.30, 0.25, 0.15, 0.20, 0.10]

# Sample from valid sales orders
sampled_sales = df_sales.sample(n=num_returns, replace=False).to_dict("records")

for s_order in sampled_sales:
    rid = f"RET{ret_counter:04d}"
    ret_counter += 1
    
    ord_date_dt = datetime.datetime.strptime(s_order["Order_Date"], "%Y-%m-%d").date()
    ret_date_dt = ord_date_dt + datetime.timedelta(days=random.randint(1, 12))
    
    qty_ret = random.randint(1, s_order["Quantity"])
    
    unit_eff_price = s_order["Sales_Amount"] / s_order["Quantity"]
    refund_amt = round(qty_ret * unit_eff_price, 2)
    
    reason = np.random.choice(return_reasons, p=reason_weights)
    
    returns_data.append({
        "Return_ID": rid,
        "Order_ID": s_order["Order_ID"],
        "Return_Date": ret_date_dt.strftime("%Y-%m-%d"),
        "Product_ID": s_order["Product_ID"],
        "Store_ID": s_order["Store_ID"],
        "Quantity_Returned": qty_ret,
        "Refund_Amount": refund_amt,
        "Return_Reason": reason
    })

df_returns = pd.DataFrame(returns_data)
print(f"Generated {len(df_returns)} returns.")

# ---------------------------------------------------------
# SAVE TO FORMATTED EXCEL WORKBOOK
# ---------------------------------------------------------
output_path = "Supply_Chain_Inventory_Analytics/data/walmart_style_supermarket_supply_chain_dataset.xlsx"
os.makedirs(os.path.dirname(output_path), exist_ok=True)

print(f"Writing to Excel workbook at: {output_path}...")

with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    df_products.to_excel(writer, sheet_name="Products", index=False)
    df_suppliers.to_excel(writer, sheet_name="Suppliers", index=False)
    df_stores.to_excel(writer, sheet_name="Stores", index=False)
    df_customers.to_excel(writer, sheet_name="Customers", index=False)
    df_inventory.to_excel(writer, sheet_name="Inventory", index=False)
    df_sales.to_excel(writer, sheet_name="Sales_Orders", index=False)
    df_purchases.to_excel(writer, sheet_name="Purchases", index=False)
    df_returns.to_excel(writer, sheet_name="Returns", index=False)

print("Applying professional formatting to Excel workbook...")

wb = openpyxl.load_workbook(output_path)

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for ws_name in wb.sheetnames:
    ws = wb[ws_name]
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    
    # Format headers
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
            # Add subtle cell border
            cell.border = thin_border
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(output_path)
print("Excel dataset generation completed successfully!")
